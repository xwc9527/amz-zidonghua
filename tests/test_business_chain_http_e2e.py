"""实机 HTTP 业务闭环：独立 uvicorn 进程 + 临时 SQLite/缓存/导出目录。"""
from __future__ import annotations

import json
import os
import socket
import sqlite3
import subprocess
import sys
import tempfile
import time
import unittest
import urllib.error
import urllib.request
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]


def _free_port() -> int:
    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0))
        return int(sock.getsockname()[1])


def _request(base: str, path: str, *, method: str = "GET", body=None):
    data = None
    headers = {}
    if body is not None:
        data = json.dumps(body).encode("utf-8")
        headers["Content-Type"] = "application/json"
    req = urllib.request.Request(base + path, data=data, headers=headers, method=method)
    try:
        with urllib.request.urlopen(req, timeout=10) as resp:
            return resp.status, json.loads(resp.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        return exc.code, json.loads(exc.read().decode("utf-8"))


class TestBusinessChainHttpE2E(unittest.TestCase):
    def test_cache_favorite_export_unfavorite_stale_and_migration_gate(self):
        with tempfile.TemporaryDirectory() as tmp:
            db = os.path.join(tmp, "categories_test.db")
            cache = os.path.join(tmp, "run_cache_test.db")
            exports = os.path.join(tmp, "exports")
            migration_lock = os.path.join(tmp, "crawl_migration.lock")
            env = os.environ.copy()
            env.update({
                "TESTING": "1",
                "DB_BACKEND": "sqlite",
                "PRODUCT_RESULT_MODE": "run_cache",
                "AMZ_DB_FILE": db,
                "DB_FILE": db,
                "AMZ_RUN_CACHE_FILE": cache,
                "AMZ_TEST_EXPORT_DIR": exports,
                "AMZ_CRAWL_MIGRATION_LOCK_FILE": migration_lock,
                "PYTHONUNBUFFERED": "1",
            })
            formal = sqlite3.connect(db)
            formal.executescript("""
                CREATE TABLE product_sightings (
                    id INTEGER PRIMARY KEY AUTOINCREMENT, asin TEXT, site TEXT
                );
                CREATE TABLE new_arrivals (
                    id INTEGER PRIMARY KEY AUTOINCREMENT, asin TEXT, site TEXT
                );
            """)
            formal.commit()
            formal.close()

            seed = subprocess.run(
                [
                    sys.executable,
                    "-c",
                    (
                        "import product_run_cache as r;"
                        "r.create_generation('HTTP-RUN-1','products');"
                        "r.activate_generation('HTTP-RUN-1','products');"
                        "r.upsert_products([{'asin':'B0HTTP0001','name':'HTTP item',"
                        "'price':19.9,'site':'US','node_id':'n1',"
                        "'list_type':'bestsellers','detail_scraped':1}],"
                        "run_id='HTTP-RUN-1',chart='products')"
                    ),
                ],
                cwd=ROOT,
                env=env,
                capture_output=True,
                text=True,
                timeout=30,
            )
            self.assertEqual(seed.returncode, 0, seed.stderr or seed.stdout)

            port = _free_port()
            base = f"http://127.0.0.1:{port}"
            server = subprocess.Popen(
                [
                    sys.executable, "-m", "uvicorn", "api_server:app",
                    "--host", "127.0.0.1", "--port", str(port), "--log-level", "warning",
                ],
                cwd=ROOT,
                env=env,
                stdout=subprocess.PIPE,
                stderr=subprocess.STDOUT,
                text=True,
            )
            lock_holder = None
            try:
                deadline = time.time() + 20
                while True:
                    if server.poll() is not None:
                        output = server.stdout.read() if server.stdout else ""
                        self.fail(f"uvicorn exited early ({server.returncode}): {output}")
                    try:
                        status, payload = _request(base, "/api/v2/favorites/count?site=US")
                        if status == 200:
                            self.assertEqual(payload["count"], 0)
                            break
                    except OSError:
                        pass
                    if time.time() >= deadline:
                        self.fail("uvicorn did not become ready within 20 seconds")
                    time.sleep(0.1)

                status, products = _request(
                    base, "/api/v2/products?site=US&detail_only=true&limit=10"
                )
                self.assertEqual(status, 200)
                self.assertIsInstance(products, list)
                self.assertEqual(len(products), 1)
                self.assertEqual(products[0]["asin"], "B0HTTP0001")
                self.assertFalse(products[0]["is_favorite"])
                cache_id = products[0]["cache_id"]
                run_id = products[0]["run_id"]

                status, added = _request(
                    base,
                    "/api/v2/favorites",
                    method="POST",
                    body={"cache_id": cache_id, "run_id": run_id},
                )
                self.assertEqual(status, 200)
                self.assertEqual(added["status"], "ok")
                self.assertEqual(added["favorite"]["asin"], "B0HTTP0001")
                formal = sqlite3.connect(db)
                try:
                    self.assertEqual(
                        formal.execute("SELECT COUNT(*) FROM product_sightings").fetchone()[0], 0
                    )
                    self.assertEqual(
                        formal.execute("SELECT COUNT(*) FROM new_arrivals").fetchone()[0], 0
                    )
                    self.assertEqual(
                        formal.execute("SELECT COUNT(*) FROM favorite_products").fetchone()[0], 1
                    )
                finally:
                    formal.close()

                status, count = _request(base, "/api/v2/favorites/count?site=US")
                self.assertEqual((status, count["count"]), (200, 1))
                status, favorites = _request(base, "/api/v2/favorites?site=US&limit=100")
                self.assertEqual(status, 200)
                self.assertEqual([row["asin"] for row in favorites], ["B0HTTP0001"])

                status, products = _request(
                    base, "/api/v2/products?site=US&detail_only=true&limit=10"
                )
                self.assertEqual(status, 200)
                self.assertTrue(products[0]["is_favorite"])

                status, exported = _request(
                    base,
                    "/api/v2/export_excel",
                    method="POST",
                    body={"source": "favorites", "site": "US"},
                )
                self.assertEqual(status, 200)
                self.assertEqual(exported["status"], "ok")
                export_file = os.path.join(exports, "favorites.xlsx")
                self.assertTrue(os.path.isfile(export_file))
                self.assertGreater(os.path.getsize(export_file), 100)
                from openpyxl import load_workbook
                book = load_workbook(export_file, read_only=True)
                try:
                    sheet = book.active
                    values = list(sheet.iter_rows(values_only=True))
                    headers = list(values[0])
                    asin_index = headers.index("asin")
                    self.assertEqual(values[1][asin_index], "B0HTTP0001")
                finally:
                    book.close()

                # 换代后旧 cache_id 必须稳定返回 409。
                switch = subprocess.run(
                    [
                        sys.executable,
                        "-c",
                        (
                            "import product_run_cache as r;"
                            "r.create_generation('HTTP-RUN-2','products');"
                            "r.activate_generation('HTTP-RUN-2','products')"
                        ),
                    ],
                    cwd=ROOT,
                    env=env,
                    capture_output=True,
                    text=True,
                    timeout=30,
                )
                self.assertEqual(switch.returncode, 0, switch.stderr or switch.stdout)
                status, stale = _request(
                    base,
                    "/api/v2/favorites",
                    method="POST",
                    body={"cache_id": cache_id, "run_id": run_id},
                )
                self.assertEqual(status, 409)
                self.assertEqual(stale["error_code"], "STALE_CACHE_ITEM")

                # 真正持有跨进程锁时，HTTP 启动接口必须在代理准备前返回 409。
                lock_holder = subprocess.Popen(
                    [
                        sys.executable,
                        "-c",
                        (
                            "import time;"
                            "from migrate_clear_crawl_results import MigrationLock;"
                            "l=MigrationLock();"
                            "assert l.acquire(0);"
                            "print('LOCK_READY',flush=True);"
                            "time.sleep(15)"
                        ),
                    ],
                    cwd=ROOT,
                    env=env,
                    stdout=subprocess.PIPE,
                    stderr=subprocess.STDOUT,
                    text=True,
                )
                self.assertEqual(lock_holder.stdout.readline().strip(), "LOCK_READY")
                status, blocked = _request(
                    base,
                    "/api/v2/start_products",
                    method="POST",
                    body={
                        "chart": "nr",
                        "site": "US",
                        "slugs": ["sandbox"],
                        "max_pages": 1,
                    },
                )
                self.assertEqual(status, 409)
                self.assertEqual(blocked["error_code"], "MIGRATION_IN_PROGRESS")

                status, deleted = _request(
                    base, "/api/v2/favorites/US/B0HTTP0001", method="DELETE"
                )
                self.assertEqual(status, 200)
                self.assertTrue(deleted["deleted"])
                status, count = _request(base, "/api/v2/favorites/count?site=US")
                self.assertEqual((status, count["count"]), (200, 0))
            finally:
                if lock_holder is not None and lock_holder.poll() is None:
                    lock_holder.terminate()
                    lock_holder.wait(timeout=5)
                if server.poll() is None:
                    server.terminate()
                    try:
                        server.wait(timeout=10)
                    except subprocess.TimeoutExpired:
                        server.kill()
                        server.wait(timeout=5)


if __name__ == "__main__":
    unittest.main()
