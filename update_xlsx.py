import sys, shutil
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
from openpyxl import load_workbook
from copy import copy

src = r"C:\Users\47763\Desktop\Amazon商品字段映射表.xlsx"
shutil.copy2(src, src.replace(".xlsx", "_backup.xlsx"))

wb = load_workbook(src)
ws = wb["商品详情页字段映射"]

def get_row_style(ws, row_num):
    styles = []
    for cell in ws[row_num]:
        styles.append({
            "font": copy(cell.font),
            "fill": copy(cell.fill),
            "alignment": copy(cell.alignment),
            "border": copy(cell.border),
        })
    return styles

def apply_style(ws, row_num, styles):
    for i, cell in enumerate(ws[row_num]):
        if i < len(styles):
            cell.font = styles[i]["font"]
            cell.fill = styles[i]["fill"]
            cell.alignment = styles[i]["alignment"]
            cell.border = styles[i]["border"]

data_style = get_row_style(ws, 32)
group_style = get_row_style(ws, 31)

def find_row(field_name):
    for row in range(1, ws.max_row + 1):
        if ws.cell(row=row, column=2).value == field_name:
            return row
    return None

def find_group_row(group_name):
    for row in range(1, ws.max_row + 1):
        if ws.cell(row=row, column=1).value == group_name:
            return row
    return None

def insert_fields(after_field, fields):
    target = find_row(after_field)
    if not target:
        print(f"  [WARN] field {after_field} not found")
        return
    insert_at = target + 1
    ws.insert_rows(insert_at, len(fields))
    for i, (f, c, sel, ex, freq, st) in enumerate(fields):
        r = insert_at + i
        ws.cell(row=r, column=2, value=f)
        ws.cell(row=r, column=3, value=c)
        ws.cell(row=r, column=4, value=sel)
        ws.cell(row=r, column=5, value=ex)
        ws.cell(row=r, column=6, value=freq)
        ws.cell(row=r, column=7, value=st)
        apply_style(ws, r, data_style)
    print(f"  Inserted {len(fields)} fields after {after_field} at row {insert_at}")

def insert_section(before_group, section_name, fields):
    target = find_group_row(before_group)
    if not target:
        print(f"  [WARN] group {before_group} not found, appending at end")
        target = ws.max_row + 1
    ws.insert_rows(target, len(fields) + 1)
    ws.cell(row=target, column=1, value=section_name)
    ws.cell(row=target, column=2, value=fields[0][0])
    ws.cell(row=target, column=3, value=fields[0][1])
    ws.cell(row=target, column=4, value=fields[0][2])
    ws.cell(row=target, column=5, value=fields[0][3])
    ws.cell(row=target, column=6, value=fields[0][4])
    ws.cell(row=target, column=7, value=fields[0][5])
    apply_style(ws, target, group_style)
    for i, (f, c, sel, ex, freq, st) in enumerate(fields[1:], 1):
        r = target + i
        ws.cell(row=r, column=2, value=f)
        ws.cell(row=r, column=3, value=c)
        ws.cell(row=r, column=4, value=sel)
        ws.cell(row=r, column=5, value=ex)
        ws.cell(row=r, column=6, value=freq)
        ws.cell(row=r, column=7, value=st)
        apply_style(ws, r, data_style)
    print(f"  Inserted section [{section_name}] with {len(fields)} fields at row {target}")

# ── 1. 购买与配送: 在 new_used_offers 后插入 ──
print("1. 购买与配送 - 新增字段")
insert_fields("new_used_offers", [
    ("shipping_fee",      "运费",         "从#deliveryBlockMessage正则提取金额", "FREE / 14,77 €", "3/6", "TEXT"),
    ("fulfillment_type",  "配送模式",     "从merchant_info推断(含Amazon配送=FBA,否则FBM)", "FBA", "3/6", "TEXT"),
    ("merchant_info",     "卖家配送方",   "#merchant-info / #merchantInfoFeature / .offer-display-feature-text", "Sold by X and Fulfilled by Amazon", "3/6", "TEXT"),
    ("has_add_to_cart",   "有加购按钮",   "#add-to-cart-button", "True", "3/6", "INT 0/1"),
    ("has_buy_now",       "有立即购买",   "#buy-now-button", "True", "3/6", "INT 0/1"),
])

# ── 2. 新增整组: 退货与保障 (插在营销标记之前) ──
print("2. 新增分组 - 退货与保障")
insert_section("营销标记", "退货与保障", [
    ("return_policy",      "退货政策", "#returnsInfoFeature_feature_div", "FREE 30-day refund / Retournierbar 30 Tagen", "3/6", "TEXT"),
    ("secure_transaction", "安全交易", "#dynamicSecureTransactionFeature_feature_div", "Your transaction is secure", "3/6", "INT 0/1"),
    ("package_info",       "包装信息", "#dynamicPackageInfoFeature_feature_div", "Ships in product packaging", "2/6", "TEXT"),
])

# ── 3. 营销标记: 在 has_sponsored_section 后插入 ──
print("3. 营销标记 - 新增字段")
insert_fields("has_sponsored_section", [
    ("social_proof",        "社交证明(月销量)", "#socialProofingAsinFaceout_feature_div", "20K+ bought in past month", "3/6", "TEXT"),
    ("bestseller_badge_text","畅销标记文本", "#zeitgeistBadge_feature_div .get_text()", "Bestseller Nr. 1 in Slides", "2/6", "TEXT"),
    ("brand_snapshot",      "品牌快照",     "#brandSnapshot_feature_div", "Top-Marke / 93% positive / 100K+", "2/6", "TEXT"),
    ("has_brand_story",     "有品牌故事",   "#aplusBrandStory_feature_div", "True", "2/6", "INT 0/1"),
    ("promotion_text",      "促销文本",     "#promoPriceBlockMessage .get_text()", "Spare 5% bei 4 Artikeln", "1/6", "TEXT"),
    ("frequently_returned", "高退货率警告", "#dpFrequentlyReturnedMessage", "Frequently returned item", "0/6", "INT 0/1"),
])

# ── 4. 媒体: 在 has_video 后插入 ──
print("4. 媒体 - 新增字段")
insert_fields("has_video", [
    ("has_3d_view",         "有3D/AR视图",   "#rhapsodyARIngress_feature_div", "True", "1/6", "INT 0/1"),
    ("has_product_videos",  "有产品视频区",  "#ive-videos-for-this-product-widget_feature_div", "True", "1/6", "INT 0/1"),
])

# ── 5. 变体: 在 variant_option_count 后插入 ──
print("5. 变体 - 新增字段")
insert_fields("variant_option_count", [
    ("has_size_chart", "有尺码表", "#sizeChartV2Data_feature_div", "True", "1/6", "INT 0/1"),
])

# ── 6. 新增整组: FOD与库存 (插在动态属性兜底之后 = 末尾) ──
print("6. 新增分组 - FOD与库存")
end_row = ws.max_row + 1
ws.cell(row=end_row, column=1, value="FOD与库存")
ws.cell(row=end_row, column=2, value="fod_message")
ws.cell(row=end_row, column=3, value="Featured Offer信息")
ws.cell(row=end_row, column=4, value="#fodcx_feature_div")
ws.cell(row=end_row, column=5, value="Keine hervorgehobenen Angebote")
ws.cell(row=end_row, column=6, value="1/6")
ws.cell(row=end_row, column=7, value="TEXT")
apply_style(ws, end_row, group_style)

for f, c, sel, ex, freq, st in [
    ("no_featured_offer", "无推荐报价", "fod_message关键词检测", "True", "1/6", "INT 0/1"),
    ("out_of_stock_msg",  "缺货信息",   "#outOfStockBuyBox_feature_div", "Dieser Artikel kann nicht...", "1/6", "TEXT"),
]:
    end_row += 1
    ws.cell(row=end_row, column=2, value=f)
    ws.cell(row=end_row, column=3, value=c)
    ws.cell(row=end_row, column=4, value=sel)
    ws.cell(row=end_row, column=5, value=ex)
    ws.cell(row=end_row, column=6, value=freq)
    ws.cell(row=end_row, column=7, value=st)
    apply_style(ws, end_row, data_style)
print(f"  Appended FOD section at rows {end_row-2}-{end_row}")

# ── 7. 修正已有字段的选择器 ──
print("7. 修正已有字段选择器")
for row in range(1, ws.max_row + 1):
    b = ws.cell(row=row, column=2).value
    if b == "delivery_info":
        ws.cell(row=row, column=4, value="#deliveryBlockMessage / #deliveryBlock_feature_div")
        ws.cell(row=row, column=5, value="FREE delivery Friday / Lieferung fuer 14,77 EUR")
        print(f"  Fixed delivery_info at row {row}")
    elif b == "buybox_info":
        ws.cell(row=row, column=4, value="#tabular-buybox / #offerDisplayFeatures_desktop")
        print(f"  Fixed buybox_info at row {row}")

# ── 8. 更新说明 ──
ws2 = wb["说明"]
total_fields = ws.max_row - 1
ws2["B7"] = f"{total_fields}个固定列 + 1个JSON兜底列"
ws2["B11"] = "2026-06-27"

wb.save(src)
print(f"\n=== 完成 === 总行数: {ws.max_row} (含表头)")

print("\n=== 全表验证 ===")
wb2 = load_workbook(src)
ws2 = wb2["商品详情页字段映射"]
current_group = ""
for row in range(1, ws2.max_row + 1):
    a = ws2.cell(row=row, column=1).value
    b = ws2.cell(row=row, column=2).value
    c = ws2.cell(row=row, column=3).value
    if a:
        current_group = a
    if b and row > 1:
        print(f"  {current_group:12s} | {b:30s} | {c}")
